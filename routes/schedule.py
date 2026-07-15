import urllib.parse
import io
import threading
from flask import Blueprint, render_template, request, redirect, url_for, flash, make_response, current_app
from flask_login import login_required, current_user
from models import db, Student, Schedule, Attendance
from datetime import date, datetime, timedelta

schedule_bp = Blueprint('schedule', __name__)

def get_week_days():
    today = date.today()
    start = today - timedelta(days=today.weekday())  # Monday
    days = []
    for i in range(7):
        d = start + timedelta(days=i)
        days.append({
            'date': d,
            'day_abbr': d.strftime('%a').upper(),
            'day_num': d.day,
            'day_name': d.strftime('%A').lower(),
            'is_today': d == today,
            'full_date': d.strftime('%A, %B %d')
        })
    return days

@schedule_bp.route('/schedule')
@login_required
def schedule():
    selected_day = request.args.get('day', date.today().strftime('%A').lower())
    week_days = get_week_days()

    schedules = Schedule.query.filter_by(
        tutor_id=current_user.id, day_of_week=selected_day, status='active'
    ).order_by(Schedule.start_time).all()

    # Get attendance for today for each schedule
    today = date.today()
    schedule_data = []
    for s in schedules:
        student = Student.query.get(s.student_id)
        if student:
            # Check if attendance already marked for today
            attendance = Attendance.query.filter_by(
                schedule_id=s.id, date=today
            ).first()
            schedule_data.append({
                'schedule': s,
                'student': student,
                'initials': ''.join([w[0].upper() for w in student.student_name.split()[:2]]),
                'attendance': attendance
            })

    students = Student.query.filter_by(tutor_id=current_user.id, status='active').order_by(Student.student_name).all()

    selected_info = next((d for d in week_days if d['day_name'] == selected_day), week_days[0])

    return render_template('schedule.html',
        week_days=week_days, schedule_data=schedule_data,
        selected_day=selected_day, selected_info=selected_info,
        students=students)

@schedule_bp.route('/schedule/add', methods=['POST'])
@login_required
def add_schedule():
    student_id = request.form.get('student_id')
    day_of_week = request.form.get('day_of_week')
    start_time = request.form.get('start_time')
    end_time = request.form.get('end_time', '')
    class_type = request.form.get('class_type', 'offline')
    location = request.form.get('location', '')

    if not student_id or not day_of_week or not start_time:
        flash('Please fill in all required fields.', 'error')
        return redirect(url_for('schedule.schedule'))

    # Validate time format
    import re
    if not re.match(r'^\d{2}:\d{2}$', start_time):
        flash('Invalid start time format.', 'error')
        return redirect(url_for('schedule.schedule'))

    sched = Schedule(
        tutor_id=current_user.id,
        student_id=int(student_id),
        day_of_week=day_of_week,
        start_time=start_time,
        end_time=end_time,
        class_type=class_type,
        location=location,
    )
    db.session.add(sched)
    db.session.commit()
    flash('Class scheduled successfully!', 'success')
    return redirect(url_for('schedule.schedule', day=day_of_week))

@schedule_bp.route('/schedule/<int:id>/delete', methods=['POST'])
@login_required
def delete_schedule(id):
    sched = Schedule.query.filter_by(id=id, tutor_id=current_user.id).first_or_404()
    day = sched.day_of_week
    db.session.delete(sched)
    db.session.commit()
    flash('Class removed from schedule.', 'success')
    return redirect(url_for('schedule.schedule', day=day))


@schedule_bp.route('/schedule/<int:id>/attendance', methods=['POST'])
@login_required
def mark_attendance(id):
    """Mark attendance for a class (completed/absent/cancelled)."""
    sched = Schedule.query.filter_by(id=id, tutor_id=current_user.id).first_or_404()
    status = request.form.get('attendance_status', 'completed')
    today = date.today()

    # Check if already marked
    existing = Attendance.query.filter_by(schedule_id=id, date=today).first()
    if existing:
        existing.status = status
    else:
        att = Attendance(
            tutor_id=current_user.id,
            student_id=sched.student_id,
            schedule_id=id,
            date=today,
            status=status
        )
        db.session.add(att)
    db.session.commit()

    labels = {'completed': 'Present', 'absent': 'Absent', 'cancelled': 'Cancelled', 'rescheduled': 'Rescheduled'}
    flash(f'Attendance marked as {labels.get(status, status)}.', 'success')
    return redirect(url_for('schedule.schedule', day=sched.day_of_week))


@schedule_bp.route('/schedule/<int:id>/remind')
@login_required
def remind_class(id):
    """Send a WhatsApp reminder for an upcoming class."""
    sched = Schedule.query.filter_by(id=id, tutor_id=current_user.id).first_or_404()
    student = Student.query.get(sched.student_id)
    if not student:
        flash('Student not found.', 'error')
        return redirect(url_for('schedule.schedule'))

    today_name = date.today().strftime('%A')
    time_str = sched.start_time
    if sched.end_time:
        time_str += f' - {sched.end_time}'

    location_str = ''
    if sched.location:
        location_str = f'\n  *Location:* {sched.location}'

    class_type_str = 'Online' if sched.class_type == 'online' else 'Offline'

    message = (
        f"------------------------------\n"
        f"       *TUITIONPE*\n"
        f"------------------------------\n\n"
        f"Namaste {student.parent_name or 'Sir/Madam'},\n\n"
        f"*Class Reminder* from *{current_user.name}*\n\n"
        f"  *Student:* {student.student_name}\n"
        f"  *Subject:* {student.subject or 'Tuition'}\n"
        f"  *Time:* {time_str}\n"
        f"  *Day:* {today_name}\n"
        f"  *Mode:* {class_type_str}{location_str}\n\n"
        f"Please ensure {student.student_name} is ready on time.\n\n"
        f"Thank you!\n"
        f"-- *{current_user.name}*\n\n"
        f"------------------------------\n"
        f"_Powered by *TuitionPe*_\n"
        f"_Smart Tuition Management_"
    )

    phone = student.parent_phone.replace('+91', '').replace(' ', '').replace('-', '').strip()
    encoded = urllib.parse.quote(message)
    link = f"https://wa.me/91{phone}?text={encoded}"

    return redirect(link)


@schedule_bp.route('/schedule/<int:id>/remind-email')
@login_required
def remind_class_email(id):
    """Send an email reminder with Google Calendar link for an online class."""
    sched = Schedule.query.filter_by(id=id, tutor_id=current_user.id).first_or_404()
    student = Student.query.get(sched.student_id)
    if not student:
        flash('Student not found.', 'error')
        return redirect(url_for('schedule.schedule'))

    if not student.parent_email:
        flash(f'No email set for {student.student_name}. Please update student details.', 'error')
        return redirect(url_for('schedule.schedule', day=sched.day_of_week))

    today = date.today()
    today_name = today.strftime('%A')

    from reminder import make_gcal_link, send_email

    time_str = sched.start_time
    if sched.end_time:
        time_str += f' - {sched.end_time}'

    gcal_link = make_gcal_link(
        title=f"{student.student_name} - {student.subject or 'Tuition'} class with {current_user.name}",
        date_obj=today,
        start_time_str=sched.start_time,
        end_time_str=sched.end_time or '',
        details=f"{student.subject or 'Tuition'} class ({sched.class_type}) with {current_user.name}",
        location=sched.location or ''
    )

    class_type_str = 'Online' if sched.class_type == 'online' else 'Offline'
    location_html = ''
    if sched.location:
        location_html = f'<tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Location</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{sched.location}</td></tr>'

    gcal_btn = ''
    if gcal_link:
        gcal_btn = f'''
        <div style="text-align:center;margin-top:20px;">
          <a href="{gcal_link}" target="_blank"
             style="display:inline-block;background:#4285f4;color:white;
                    border-radius:10px;padding:14px 28px;font-size:14px;font-weight:600;
                    text-decoration:none;">
            &#128197; Add to Google Calendar
          </a>
        </div>'''

    html = f'''<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
<body style="margin:0;padding:0;background:#050505;font-family:'Segoe UI',Arial,sans-serif;">
  <div style="max-width:480px;margin:0 auto;padding:24px 12px;">
    <div style="text-align:center;margin-bottom:20px;">
      <span style="color:#10b981;font-size:16px;font-weight:800;letter-spacing:5px;">TUITIONPE</span>
    </div>
    <div style="background:#0a0a0a;border:1px solid rgba(255,255,255,0.07);border-radius:16px;overflow:hidden;">
      <div style="background:linear-gradient(135deg,#10b981,#059669);padding:24px;text-align:center;">
        <p style="color:rgba(255,255,255,0.85);margin:0 0 6px;font-size:11px;text-transform:uppercase;letter-spacing:2px;">Class Reminder</p>
        <h2 style="color:white;margin:0;font-size:20px;font-weight:700;">Upcoming Class Today</h2>
      </div>
      <div style="padding:24px;">
        <p style="color:#9ca3af;font-size:13px;margin:0 0 16px;">Namaste {student.parent_name or 'Sir/Madam'},</p>
        <p style="color:white;font-size:14px;margin:0 0 20px;">This is a reminder from <strong>{current_user.name}</strong> about today's class:</p>
        <table style="width:100%;border-collapse:collapse;background:rgba(255,255,255,0.03);border-radius:10px;overflow:hidden;">
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Student</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{student.student_name}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Subject</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{student.subject or 'Tuition'}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Time</td><td style="color:#10b981;padding:6px 12px;font-size:13px;font-weight:700;">{time_str}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Day</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{today_name}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Mode</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{class_type_str}</td></tr>
          {location_html}
        </table>
        {gcal_btn}
        <p style="color:#9ca3af;font-size:13px;margin:20px 0 0;">Please ensure {student.student_name} is ready on time.</p>
        <p style="color:#9ca3af;font-size:13px;margin:8px 0 0;">Thank you!<br><strong style="color:white;">{current_user.name}</strong></p>
      </div>
    </div>
    <div style="text-align:center;margin-top:20px;">
      <span style="color:#10b981;font-size:11px;font-weight:800;letter-spacing:3px;">TUITIONPE</span>
      <p style="color:#374151;font-size:10px;margin:4px 0 0;">Smart Tuition Management for Modern Teachers</p>
    </div>
  </div>
</body>
</html>'''

    subject = f"Class Reminder - {student.student_name} ({student.subject or 'Tuition'}) at {sched.start_time}"
    app = current_app._get_current_object()
    parent_email = student.parent_email

    def _send_bg():
        from reminder import send_email as rem_send_email
        ok = rem_send_email(app, parent_email, subject, html)
        print(f"[TuitionPe] Class email reminder {'sent' if ok else 'FAILED'} -> {parent_email}")

    t = threading.Thread(target=_send_bg, daemon=True)
    t.start()

    flash(f'Email reminder with Google Calendar link sent to {parent_email}!', 'success')
    return redirect(url_for('schedule.schedule', day=sched.day_of_week))


@schedule_bp.route('/schedule/remind-all-email')
@login_required
def remind_all_email():
    """Send email reminders to ALL today's students who have parent_email set."""
    today_day = date.today().strftime('%A').lower()
    schedules = Schedule.query.filter_by(
        tutor_id=current_user.id, day_of_week=today_day, status='active'
    ).order_by(Schedule.start_time).all()

    if not schedules:
        flash('No classes scheduled for today.', 'error')
        return redirect(url_for('schedule.schedule'))

    from reminder import make_gcal_link, send_email as rem_send_email
    today = date.today()
    today_name = today.strftime('%A')
    app = current_app._get_current_object()
    sent_count = 0
    skip_count = 0

    for sched in schedules:
        student = Student.query.get(sched.student_id)
        if not student or not student.parent_email:
            skip_count += 1
            continue

        time_str = sched.start_time
        if sched.end_time:
            time_str += f' - {sched.end_time}'

        gcal_link = make_gcal_link(
            title=f"{student.student_name} - {student.subject or 'Tuition'} class with {current_user.name}",
            date_obj=today,
            start_time_str=sched.start_time,
            end_time_str=sched.end_time or '',
            details=f"{student.subject or 'Tuition'} class ({sched.class_type}) with {current_user.name}",
            location=sched.location or ''
        )

        class_type_str = 'Online' if sched.class_type == 'online' else 'Offline'
        location_html = ''
        if sched.location:
            location_html = f'<tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Location</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{sched.location}</td></tr>'

        gcal_btn = ''
        if gcal_link:
            gcal_btn = f'''
            <div style="text-align:center;margin-top:20px;">
              <a href="{gcal_link}" target="_blank"
                 style="display:inline-block;background:#4285f4;color:white;
                        border-radius:10px;padding:14px 28px;font-size:14px;font-weight:600;
                        text-decoration:none;">
                &#128197; Add to Google Calendar
              </a>
            </div>'''

        html = f'''<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
<body style="margin:0;padding:0;background:#050505;font-family:'Segoe UI',Arial,sans-serif;">
  <div style="max-width:480px;margin:0 auto;padding:24px 12px;">
    <div style="text-align:center;margin-bottom:20px;">
      <span style="color:#10b981;font-size:16px;font-weight:800;letter-spacing:5px;">TUITIONPE</span>
    </div>
    <div style="background:#0a0a0a;border:1px solid rgba(255,255,255,0.07);border-radius:16px;overflow:hidden;">
      <div style="background:linear-gradient(135deg,#10b981,#059669);padding:24px;text-align:center;">
        <p style="color:rgba(255,255,255,0.85);margin:0 0 6px;font-size:11px;text-transform:uppercase;letter-spacing:2px;">Class Reminder</p>
        <h2 style="color:white;margin:0;font-size:20px;font-weight:700;">Upcoming Class Today</h2>
      </div>
      <div style="padding:24px;">
        <p style="color:#9ca3af;font-size:13px;margin:0 0 16px;">Namaste {student.parent_name or 'Sir/Madam'},</p>
        <p style="color:white;font-size:14px;margin:0 0 20px;">This is a reminder from <strong>{current_user.name}</strong> about today's class:</p>
        <table style="width:100%;border-collapse:collapse;background:rgba(255,255,255,0.03);border-radius:10px;overflow:hidden;">
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Student</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{student.student_name}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Subject</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{student.subject or 'Tuition'}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Time</td><td style="color:#10b981;padding:6px 12px;font-size:13px;font-weight:700;">{time_str}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Day</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{today_name}</td></tr>
          <tr><td style="color:#9ca3af;padding:6px 12px;font-size:13px;">Mode</td><td style="color:white;padding:6px 12px;font-size:13px;font-weight:600;">{class_type_str}</td></tr>
          {location_html}
        </table>
        {gcal_btn}
        <p style="color:#9ca3af;font-size:13px;margin:20px 0 0;">Please ensure {student.student_name} is ready on time.</p>
        <p style="color:#9ca3af;font-size:13px;margin:8px 0 0;">Thank you!<br><strong style="color:white;">{current_user.name}</strong></p>
      </div>
    </div>
    <div style="text-align:center;margin-top:20px;">
      <span style="color:#10b981;font-size:11px;font-weight:800;letter-spacing:3px;">TUITIONPE</span>
      <p style="color:#374151;font-size:10px;margin:4px 0 0;">Smart Tuition Management for Modern Teachers</p>
    </div>
  </div>
</body>
</html>'''

        subject = f"Class Reminder - {student.student_name} ({student.subject or 'Tuition'}) at {sched.start_time}"
        parent_email = student.parent_email

        def _send_bg(app_ctx, to_email, subj, body):
            ok = rem_send_email(app_ctx, to_email, subj, body)
            print(f"[TuitionPe] Bulk email {'sent' if ok else 'FAILED'} -> {to_email}")

        t = threading.Thread(target=_send_bg, args=(app, parent_email, subject, html), daemon=True)
        t.start()
        sent_count += 1

    if sent_count > 0:
        msg = f'Email reminders sent to {sent_count} student(s)!'
        if skip_count > 0:
            msg += f' ({skip_count} skipped - no email set)'
        flash(msg, 'success')
    else:
        flash('No students have email addresses set. Please update student details.', 'error')

    return redirect(url_for('schedule.schedule'))


@schedule_bp.route('/schedule/remind-all')
@login_required
def remind_all_today():
    """Build a page with WhatsApp links for ALL today's classes instead of just the first."""
    today_day = date.today().strftime('%A').lower()
    schedules = Schedule.query.filter_by(
        tutor_id=current_user.id, day_of_week=today_day, status='active'
    ).order_by(Schedule.start_time).all()

    if not schedules:
        flash('No classes scheduled for today.', 'error')
        return redirect(url_for('schedule.schedule'))

    # Build reminder links for ALL classes
    reminder_links = []
    for sched in schedules:
        student = Student.query.get(sched.student_id)
        if not student:
            continue
        today_name = date.today().strftime('%A')
        time_str = sched.start_time
        if sched.end_time:
            time_str += f' - {sched.end_time}'
        location_str = ''
        if sched.location:
            location_str = f'\n  *Location:* {sched.location}'
        class_type_str = 'Online' if sched.class_type == 'online' else 'Offline'

        message = (
            f"------------------------------\n"
            f"       *TUITIONPE*\n"
            f"------------------------------\n\n"
            f"Namaste {student.parent_name or 'Sir/Madam'},\n\n"
            f"*Class Reminder* from *{current_user.name}*\n\n"
            f"  *Student:* {student.student_name}\n"
            f"  *Subject:* {student.subject or 'Tuition'}\n"
            f"  *Time:* {time_str}\n"
            f"  *Day:* {today_name}\n"
            f"  *Mode:* {class_type_str}{location_str}\n\n"
            f"Please ensure {student.student_name} is ready on time.\n\n"
            f"Thank you!\n"
            f"-- *{current_user.name}*\n\n"
            f"------------------------------\n"
            f"_Powered by *TuitionPe*_\n"
            f"_Smart Tuition Management_"
        )

        phone = student.parent_phone.replace('+91', '').replace(' ', '').replace('-', '').strip()
        encoded = urllib.parse.quote(message)
        link = f"https://wa.me/91{phone}?text={encoded}"

        reminder_links.append({
            'student_name': student.student_name,
            'subject': student.subject or 'Tuition',
            'time': sched.start_time,
            'link': link,
            'schedule_id': sched.id,
            'has_email': bool(student.parent_email),
            'class_type': sched.class_type,
        })

    return render_template('remind_all.html', reminder_links=reminder_links)


# ────────────────────────────────────────────────────────────
# Attendance Tracker — View, Filter, Export
# ────────────────────────────────────────────────────────────

@schedule_bp.route('/attendance')
@login_required
def attendance_tracker():
    """Attendance tracker page with filters: student, date range, status."""
    # Get filter values
    student_id = request.args.get('student_id', '', type=str)
    status_filter = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')

    # Default: last 30 days
    if not date_from:
        date_from = (date.today() - timedelta(days=30)).isoformat()
    if not date_to:
        date_to = date.today().isoformat()

    # Parse dates
    try:
        from_date = date.fromisoformat(date_from)
    except (ValueError, TypeError):
        from_date = date.today() - timedelta(days=30)
    try:
        to_date = date.fromisoformat(date_to)
    except (ValueError, TypeError):
        to_date = date.today()

    # Build query
    query = Attendance.query.filter(
        Attendance.tutor_id == current_user.id,
        Attendance.date >= from_date,
        Attendance.date <= to_date
    )

    if student_id:
        try:
            query = query.filter(Attendance.student_id == int(student_id))
        except (ValueError, TypeError):
            pass

    if status_filter and status_filter in ('completed', 'absent', 'cancelled', 'rescheduled'):
        query = query.filter(Attendance.status == status_filter)

    records = query.order_by(Attendance.date.desc(), Attendance.id.desc()).all()

    # Enrich records with student/schedule info
    attendance_data = []
    for rec in records:
        student = Student.query.get(rec.student_id)
        schedule = Schedule.query.get(rec.schedule_id)
        if student and schedule:
            attendance_data.append({
                'record': rec,
                'student': student,
                'schedule': schedule,
                'initials': ''.join([w[0].upper() for w in student.student_name.split()[:2]])
            })

    # Summary counts
    total = len(attendance_data)
    completed_count = sum(1 for a in attendance_data if a['record'].status == 'completed')
    absent_count = sum(1 for a in attendance_data if a['record'].status == 'absent')
    cancelled_count = sum(1 for a in attendance_data if a['record'].status == 'cancelled')
    rescheduled_count = sum(1 for a in attendance_data if a['record'].status == 'rescheduled')

    # Student list for filter dropdown
    students = Student.query.filter_by(tutor_id=current_user.id, status='active').order_by(Student.student_name).all()

    return render_template('attendance_tracker.html',
        attendance_data=attendance_data,
        students=students,
        student_id=student_id,
        status_filter=status_filter,
        date_from=from_date.isoformat(),
        date_to=to_date.isoformat(),
        total=total,
        completed_count=completed_count,
        absent_count=absent_count,
        cancelled_count=cancelled_count,
        rescheduled_count=rescheduled_count)


@schedule_bp.route('/attendance/export')
@login_required
def export_attendance():
    """Export attendance data as XLSX with TuitionPe branding."""
    # Get same filters as the tracker page
    student_id = request.args.get('student_id', '', type=str)
    status_filter = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')

    if not date_from:
        date_from = (date.today() - timedelta(days=30)).isoformat()
    if not date_to:
        date_to = date.today().isoformat()

    try:
        from_date = date.fromisoformat(date_from)
    except (ValueError, TypeError):
        from_date = date.today() - timedelta(days=30)
    try:
        to_date = date.fromisoformat(date_to)
    except (ValueError, TypeError):
        to_date = date.today()

    query = Attendance.query.filter(
        Attendance.tutor_id == current_user.id,
        Attendance.date >= from_date,
        Attendance.date <= to_date
    )
    if student_id:
        try:
            query = query.filter(Attendance.student_id == int(student_id))
        except (ValueError, TypeError):
            pass
    if status_filter and status_filter in ('completed', 'absent', 'cancelled', 'rescheduled'):
        query = query.filter(Attendance.status == status_filter)

    records = query.order_by(Attendance.date.desc()).all()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash('Export requires openpyxl. Run: pip install openpyxl', 'error')
        return redirect(url_for('schedule.attendance_tracker'))

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # ── Branding / Header ──
    # Merge cells for title
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = 'TUITIONPE'
    title_cell.font = Font(name='Calibri', size=22, bold=True, color='10B981')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    ws.merge_cells('A2:H2')
    subtitle_cell = ws['A2']
    subtitle_cell.value = 'Attendance Report'
    subtitle_cell.font = Font(name='Calibri', size=13, bold=True, color='374151')
    subtitle_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 25

    ws.merge_cells('A3:H3')
    info_cell = ws['A3']
    info_cell.value = f'Teacher: {current_user.name}  |  Period: {from_date.strftime("%d %b %Y")} to {to_date.strftime("%d %b %Y")}  |  Generated: {datetime.now().strftime("%d %b %Y %I:%M %p")}'
    info_cell.font = Font(name='Calibri', size=9, color='6B7280')
    info_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 20

    # Empty row
    ws.row_dimensions[4].height = 10

    # ── Column headers ──
    headers = ['#', 'Date', 'Day', 'Student Name', 'Subject', 'Class Time', 'Status', 'Notes']
    header_fill = PatternFill(start_color='10B981', end_color='10B981', fill_type='solid')
    header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[5].height = 28

    # ── Data rows ──
    status_colors = {
        'completed': '10B981',
        'absent': 'EF4444',
        'cancelled': '6B7280',
        'rescheduled': 'F59E0B'
    }
    alt_fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

    for row_idx, rec in enumerate(records, 1):
        student = Student.query.get(rec.student_id)
        schedule = Schedule.query.get(rec.schedule_id)
        if not student or not schedule:
            continue

        excel_row = row_idx + 5  # offset by header rows
        day_name = rec.date.strftime('%A')
        time_str = schedule.start_time
        if schedule.end_time:
            time_str += f' - {schedule.end_time}'

        row_data = [
            row_idx,
            rec.date.strftime('%d %b %Y'),
            day_name,
            student.student_name,
            student.subject or 'Tuition',
            time_str,
            rec.status.capitalize(),
            rec.notes or ''
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10, color='374151')
            cell.alignment = Alignment(horizontal='center' if col_idx in (1, 2, 3, 6, 7) else 'left', vertical='center')
            cell.border = thin_border

            # Alternate row shading
            if row_idx % 2 == 0:
                cell.fill = alt_fill

            # Status column coloring
            if col_idx == 7:
                status_clr = status_colors.get(rec.status, '374151')
                cell.font = Font(name='Calibri', size=10, bold=True, color=status_clr)

    # ── Column widths ──
    col_widths = [5, 14, 12, 22, 18, 16, 14, 25]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ── Summary row ──
    summary_row = len(records) + 7
    ws.merge_cells(f'A{summary_row}:H{summary_row}')
    ws.row_dimensions[summary_row].height = 10  # spacer

    summary_row += 1
    ws.merge_cells(f'A{summary_row}:H{summary_row}')
    cell = ws[f'A{summary_row}']
    completed = sum(1 for r in records if r.status == 'completed')
    absent = sum(1 for r in records if r.status == 'absent')
    cancelled = sum(1 for r in records if r.status == 'cancelled')
    rescheduled = sum(1 for r in records if r.status == 'rescheduled')
    cell.value = f'Summary:  Total: {len(records)}  |  Completed: {completed}  |  Absent: {absent}  |  Cancelled: {cancelled}  |  Rescheduled: {rescheduled}'
    cell.font = Font(name='Calibri', size=10, bold=True, color='374151')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[summary_row].height = 25

    # ── Footer / Branding ──
    footer_row = summary_row + 2
    ws.merge_cells(f'A{footer_row}:H{footer_row}')
    cell = ws[f'A{footer_row}']
    cell.value = 'TuitionPe - Smart Tuition Management for Modern Teachers'
    cell.font = Font(name='Calibri', size=9, italic=True, color='10B981')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    # ── Print header/footer watermark ──
    ws.oddHeader.center.text = 'TUITIONPE'
    ws.oddHeader.center.size = 14
    ws.oddHeader.center.color = '10B981'
    ws.oddFooter.center.text = 'TuitionPe - Smart Tuition Management'
    ws.oddFooter.center.size = 9

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"TuitionPe_Attendance_{from_date.strftime('%Y%m%d')}_{to_date.strftime('%Y%m%d')}.xlsx"
    response = make_response(buffer.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
